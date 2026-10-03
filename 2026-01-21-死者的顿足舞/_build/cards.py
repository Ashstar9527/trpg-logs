"""读六张人物卡（COC7 空白卡模板的「简化卡」页）-> _build/cards.json"""
import openpyxl, json, os, re, warnings
warnings.filterwarnings('ignore')
D = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORDER = ['尤金', '克洛德', '文森特', '蒂诺丝', '朱利安', '西奥多']


def v(ws, a):
    x = ws[a].value
    return None if x in (None, 0, '0', ' ', '') else x


out = []
for n in ORDER:
    wb = openpyxl.load_workbook(os.path.join(D, f'角色卡-{n}.xlsx'), data_only=True)
    ws, pc = wb[next(x for x in wb.sheetnames if x.startswith('简化卡'))], wb['人物卡']
    row_of = {str(ws[f'B{r}'].value).strip(): r for r in range(1, 40) if ws[f'B{r}'].value}
    R_ITEM, R_DESC = row_of['随身物品'] + 1, row_of['描述']
    c = dict(key=n, summary=v(ws, 'B2'), player=None)
    for row in pc.iter_rows(max_row=6):
        for cell in row:
            if cell.value == '玩家':
                c['player'] = next((x.value for x in row[cell.column:] if x.value), None)
    c['attr'] = {k: v(ws, a) for k, a in [('STR', 'C3'), ('DEX', 'F3'), ('POW', 'I3'), ('INT', 'L3'), ('CON', 'C5'),
                                           ('APP', 'F5'), ('EDU', 'I5'), ('SIZ', 'L5')]}
    c['derived'] = {k: v(ws, a) for k, a in [('HP', 'C7'), ('SAN', 'F7'), ('MP', 'I7'), ('LUCK', 'I9'), ('MOV', 'L9'), ('DB', 'C9')]}
    sk = []
    for r in range(3, 37):
        for nc, vc in (('N', 'Q'), ('T', 'W')):
            nm, val = v(ws, f'{nc}{r}'), v(ws, f'{vc}{r}')
            if nm and isinstance(val, (int, float)) and val:
                nm = str(nm).strip('：: ')
                sub = v(ws, f'{chr(ord(nc) + 1)}{r}')
                if sub and str(sub).strip():
                    nm = f'{nm}（{sub}）'
                sk.append((nm, int(val)))
    c['skills'] = sorted(sk, key=lambda x: -x[1])
    c['weapons'] = [dict(name=v(ws, f'B{r}') or v(ws, f'C{r}'), skill=v(ws, f'D{r}'), val=v(ws, f'E{r}'), dmg=v(ws, f'H{r}'))
                    for r in range(12, 17) if (v(ws, f'B{r}') or v(ws, f'C{r}')) and v(ws, f'B{r}') != '无' and v(ws, f'E{r}')]
    items = [v(ws, f'{col}{r}') for r in range(R_ITEM, R_DESC - 1) for col in 'BEHK']
    c['items'] = list(dict.fromkeys(i for i in items if i))
    c['bg'] = {lab: v(ws, f'C{R_DESC + k}') for k, lab in enumerate(['描述', '信仰', '重要之人', '重要之地', '宝物', '特质', '伤疤', '恐惧'])
               if v(ws, f'C{R_DESC + k}') not in (None, '无')}
    c['story'] = v(ws, f'H{R_DESC}')
    out.append(c)
json.dump(out, open(os.path.join(D, '_build', 'cards.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
for c in out:
    print('==', c['key'], c['player'], c['summary']); print('  ', c['attr'], c['derived'])
    print('  ', c['skills'][:14]); print('  ', c['weapons']); print('  ', c['items']); print('  ', c['bg']); print('   story', len(c['story'] or ''))
